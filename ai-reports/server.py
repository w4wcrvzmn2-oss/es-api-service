# -*- coding: utf-8 -*-
"""
AI-сервис отчётов для покупателей (Электронная Фармация).

Поток:
  десктоп -> POST /generate {prompt, token, format?}
    -> Ollama (локально) превращает промт в JSON-интент (тип отчёта, даты, фильтры)
    -> платформа /api/buyer/report (данные СТРОГО этого покупателя по его JWT)
    -> детерминированная агрегация + сборка Word/Excel
  -> файл возвращается клиенту

LLM только понимает формулировку. Цифры считает код -> отчёт всегда точный,
данные — только клиента (скоуп на сервере по токену).
"""
import os
import io
import json
import re
import datetime as dt
from typing import Optional, List, Dict, Any

import requests
from fastapi import FastAPI, HTTPException
from fastapi.responses import StreamingResponse, JSONResponse
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

# ---- конфигурация (можно переопределить через окружение) ----
OLLAMA_URL = os.environ.get("OLLAMA_URL", "http://127.0.0.1:11434")
OLLAMA_MODEL = os.environ.get("OLLAMA_MODEL", "qwen2.5:7b-instruct")
API_BASE = os.environ.get("ELF_API_BASE", "https://24pharmdata.ru").rstrip("/")
HTTP_TIMEOUT = int(os.environ.get("ELF_HTTP_TIMEOUT", "60"))

app = FastAPI(title="Elfisa AI Reports", version="1.0")

# CORS — чтобы кабинет phd (phd.24pharmdata.ru) мог звать /ai/chat кросс-доменно
# из браузера. На проксируемые (Caddy) POST-запросы не влияет: 400 «error parsing
# the body» ловил curl с Expect:100-continue, а не CORS — браузер/десктоп его не шлют.
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)

REPORT_TYPES = {"orders", "orders_items", "by_supplier", "by_item", "by_pharmacy"}

RU_TITLES = {
    "orders": "Мои заказы за период",
    "orders_items": "Мои заказы с товарами",
    "by_supplier": "Закупки по поставщикам",
    "by_item": "Закупки по товарам",
    "by_pharmacy": "Закупки по аптекам",
}


class GenReq(BaseModel):
    prompt: str
    token: str
    format: Optional[str] = None  # docx | xlsx (по умолчанию из интента/docx)


class MapReq(BaseModel):
    # columns: [{"name": str, "samples": [str, ...]}]
    # fields:  [{"key": str, "description": str}]
    columns: List[Dict[str, Any]]
    fields: List[Dict[str, Any]]


class ChatReq(BaseModel):
    message: str
    history: Optional[List[Dict[str, str]]] = None  # [{"role": "...", "content": "..."}]


# Персона чат-ассистента ExestAI (раздел «Чат» в кабинете).
EXEST_AI_SYSTEM = (
    "Ты — ExestAI, дружелюбный ассистент платформы «ЭльФиСА» (Электронная Фармация) — "
    "сервиса заказа лекарств для аптек и кабинета поставщика. Помогаешь пользователю: "
    "объясняешь как работать с прайсами, заказами, поставщиками, выгрузками, отчётами; "
    "отвечаешь на общие вопросы. Кратко, вежливо и по делу. "
    "Если чего-то не знаешь наверняка — честно скажи об этом. "
    "ВАЖНО: отвечай СТРОГО на русском языке. Никогда не используй китайские иероглифы "
    "или другие языки — только русский (латиница допустима лишь для названий и терминов)."
)

# Запрос «почему мы круче / сравни с конкурентом / реклама» — это НЕ отчёт по данным,
# а маркетинговый текст. Такие промты уводим в рекламный режим (честно помечаем).
MARKETING_KEYWORDS = (
    "круче", "сравни", "сравнени", "по сравнению", "реклам", "преимуществ",
    "конкурент", "аналит", "чем хорош", "чем лучш", "почему выбрать",
    "почему мы", "почему эта программа", "программа класс", "класс по",
    "лучше всех", "лучше конкурент", "достоинств",
)


def _today() -> dt.date:
    return dt.date.today()


def _default_range():
    to = _today()
    frm = (to.replace(day=1) - dt.timedelta(days=1)).replace(day=1)  # начало прошлого месяца
    return frm, to


def _parse_date(s: Any, fallback: dt.date) -> dt.date:
    if not s:
        return fallback
    try:
        return dt.datetime.strptime(str(s)[:10], "%Y-%m-%d").date()
    except Exception:
        return fallback


def extract_intent(prompt: str) -> Dict[str, Any]:
    """Промт -> JSON-интент через локальную Ollama (format=json)."""
    today = _today().isoformat()
    system = (
        "Ты помощник, который переводит запрос фармацевта в параметры отчёта. "
        "Верни СТРОГО один JSON-объект без пояснений со схемой: "
        "{\"report_type\": один из [orders, orders_items, by_supplier, by_item, by_pharmacy], "
        "\"date_from\": \"YYYY-MM-DD\", \"date_to\": \"YYYY-MM-DD\", "
        "\"supplier_filter\": строка или null, \"item_filter\": строка или null, "
        "\"format\": \"docx\" или \"xlsx\", \"title\": краткий заголовок на русском}. "
        "Значения report_type: orders=список заказов, по одной строке на заказ, БЕЗ товаров внутри; "
        "orders_items=заказы С ИХ СОДЕРЖИМЫМ — каждая позиция (товар) внутри заказа отдельной строкой; "
        "выбирай orders_items когда просят «с содержимым», «с товарами», «что внутри», «детально», «позиции заказов», «расшифровка»; "
        "by_supplier=сводка по поставщикам; by_item=по товарам/препаратам; "
        "by_pharmacy=по аптекам-грузополучателям. "
        "«за всё время» — период с 2023-01-01 по сегодня. "
        f"Сегодня {today}. Если период не указан — последний месяц. "
        "Если тип не ясен — orders. Если формат не указан — docx."
    )
    payload = {
        "model": OLLAMA_MODEL,
        "format": "json",
        "stream": False,
        "options": {"temperature": 0.1},
        "messages": [
            {"role": "system", "content": system},
            {"role": "user", "content": prompt},
        ],
    }
    r = requests.post(f"{OLLAMA_URL}/api/chat", json=payload, timeout=HTTP_TIMEOUT)
    r.raise_for_status()
    content = r.json().get("message", {}).get("content", "{}")
    try:
        intent = json.loads(content)
    except Exception:
        m = re.search(r"\{.*\}", content, re.S)
        intent = json.loads(m.group(0)) if m else {}
    return intent


def normalize_intent(intent: Dict[str, Any], req_format: Optional[str]) -> Dict[str, Any]:
    frm_def, to_def = _default_range()
    rtype = str(intent.get("report_type", "orders")).strip()
    if rtype not in REPORT_TYPES:
        rtype = "orders"
    fmt = (req_format or intent.get("format") or "docx").strip().lower()
    if fmt not in ("docx", "xlsx"):
        fmt = "docx"
    return {
        "report_type": rtype,
        "date_from": _parse_date(intent.get("date_from"), frm_def),
        "date_to": _parse_date(intent.get("date_to"), to_def),
        "supplier_filter": (intent.get("supplier_filter") or None),
        "item_filter": (intent.get("item_filter") or None),
        "format": fmt,
        "title": (intent.get("title") or RU_TITLES.get(rtype)) or "Отчёт",
    }


def fetch_lines(token: str, frm: dt.date, to: dt.date) -> List[Dict[str, Any]]:
    """Данные строго покупателя — платформа скоупит по его JWT."""
    url = f"{API_BASE}/api/buyer/report"
    params = {"date_from": frm.isoformat(), "date_to": to.isoformat()}
    headers = {"Authorization": f"Bearer {token}"}
    r = requests.get(url, params=params, headers=headers, timeout=HTTP_TIMEOUT)
    if r.status_code == 401:
        raise HTTPException(status_code=401, detail="Токен недействителен или истёк")
    r.raise_for_status()
    return r.json().get("lines", []) or []


def _num(x):
    try:
        return float(x or 0)
    except Exception:
        return 0.0


def apply_filters(lines, sup_filter, item_filter):
    out = lines
    if sup_filter:
        s = str(sup_filter).lower()
        out = [l for l in out if s in str(l.get("supplier", "")).lower()]
    if item_filter:
        s = str(item_filter).lower()
        out = [l for l in out if s in str(l.get("item_name", "")).lower()]
    return out


def aggregate(lines, rtype):
    """Возвращает (columns, rows). Логика зеркалит десктопный UCReports."""
    if rtype == "orders":
        cols = ["Дата", "Номер", "Поставщик", "Аптека", "Позиций", "Сумма", "Статус"]
        groups: Dict[str, list] = {}
        for l in lines:
            groups.setdefault(l.get("order_id", ""), []).append(l)
        rows = []
        for _, g in groups.items():
            f = g[0]
            date = (f.get("order_date") or "")[:10]
            suppliers = ", ".join(sorted({l.get("supplier", "") for l in g if l.get("supplier")}))
            rows.append([
                date, f.get("global_sign", ""), suppliers, f.get("location", ""),
                len(g), sum(_num(l.get("sum")) for l in g), f.get("status", ""),
            ])
        rows.sort(key=lambda r: r[0], reverse=True)
        return cols, rows

    if rtype == "orders_items":
        # Заказы с содержимым: каждая позиция (товар) отдельной строкой.
        cols = ["Дата", "Номер", "Поставщик", "Товар", "Код", "Кол-во", "Цена", "Сумма", "Статус"]
        rows = []
        for l in sorted(lines, key=lambda x: ((x.get("order_date") or ""), x.get("global_sign") or ""), reverse=True):
            rows.append([
                (l.get("order_date") or "")[:10], l.get("global_sign", ""), l.get("supplier", ""),
                l.get("item_name", ""), l.get("item_code", ""),
                _num(l.get("qty")), _num(l.get("unit_price")), _num(l.get("sum")), l.get("status", ""),
            ])
        return cols, rows

    if rtype == "by_supplier":
        cols = ["Поставщик", "Заказов", "Позиций", "Сумма"]
        groups = {}
        for l in lines:
            groups.setdefault(l.get("supplier") or "(не указан)", []).append(l)
        rows = [[k, len({l.get("order_id") for l in g}), len(g), sum(_num(l.get("sum")) for l in g)]
                for k, g in groups.items()]
        rows.sort(key=lambda r: r[3], reverse=True)
        return cols, rows

    if rtype == "by_item":
        cols = ["Товар", "Код", "Кол-во", "Сумма"]
        groups = {}
        for l in lines:
            key = (str(l.get("item_name", "")).strip(), str(l.get("item_code", "")))
            groups.setdefault(key, []).append(l)
        rows = [[k[0], k[1], sum(_num(l.get("qty")) for l in g), sum(_num(l.get("sum")) for l in g)]
                for k, g in groups.items()]
        rows.sort(key=lambda r: r[3], reverse=True)
        return cols, rows

    # by_pharmacy
    cols = ["Аптека", "Заказов", "Позиций", "Сумма"]
    groups = {}
    for l in lines:
        groups.setdefault(l.get("location") or "(не указана)", []).append(l)
    rows = [[k, len({l.get("order_id") for l in g}), len(g), sum(_num(l.get("sum")) for l in g)]
            for k, g in groups.items()]
    rows.sort(key=lambda r: r[3], reverse=True)
    return cols, rows


def ai_summary(title, period, cols, rows, total_sum) -> str:
    """Короткое резюме 2-3 предложения (best-effort, не критично)."""
    try:
        preview = [dict(zip(cols, r)) for r in rows[:15]]
        system = ("Ты аналитик. По данным отчёта напиши 2-3 коротких предложения-резюме "
                  "на русском: ключевые цифры, лидеры, заметные моменты. Без вступлений.")
        user = (f"Отчёт: {title}. Период: {period}. Итоговая сумма: {total_sum:.2f}. "
                f"Строки (до 15): {json.dumps(preview, ensure_ascii=False, default=str)}")
        payload = {"model": OLLAMA_MODEL, "stream": False, "options": {"temperature": 0.3},
                   "messages": [{"role": "system", "content": system},
                                {"role": "user", "content": user}]}
        r = requests.post(f"{OLLAMA_URL}/api/chat", json=payload, timeout=HTTP_TIMEOUT)
        r.raise_for_status()
        return r.json().get("message", {}).get("content", "").strip()
    except Exception:
        return ""


def _fmt_cell(col, val):
    if col in ("Сумма", "Цена"):
        return f"{_num(val):,.2f}".replace(",", " ")
    if col == "Кол-во":
        return f"{_num(val):,.3f}".replace(",", " ")
    return "" if val is None else str(val)


def build_docx(meta, cols, rows, summary) -> bytes:
    doc = Document()
    h = doc.add_heading(meta["title"], level=1)
    h.alignment = WD_ALIGN_PARAGRAPH.LEFT
    period = f'Период: {meta["date_from"].strftime("%d.%m.%Y")} — {meta["date_to"].strftime("%d.%m.%Y")}'
    p = doc.add_paragraph(period)
    p.runs[0].italic = True
    if meta.get("supplier_filter"):
        doc.add_paragraph(f'Фильтр по поставщику: {meta["supplier_filter"]}')
    if meta.get("item_filter"):
        doc.add_paragraph(f'Фильтр по товару: {meta["item_filter"]}')
    if summary:
        sp = doc.add_paragraph()
        run = sp.add_run(summary)
        run.font.size = Pt(11)

    table = doc.add_table(rows=1, cols=len(cols))
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    for i, c in enumerate(cols):
        hdr[i].text = c
        for par in hdr[i].paragraphs:
            for run in par.runs:
                run.bold = True

    total_sum = 0.0
    sum_idx = cols.index("Сумма") if "Сумма" in cols else None
    for r in rows:
        cells = table.add_row().cells
        for i, c in enumerate(cols):
            cells[i].text = _fmt_cell(c, r[i])
        if sum_idx is not None:
            total_sum += _num(r[sum_idx])

    doc.add_paragraph()
    tp = doc.add_paragraph()
    tr = tp.add_run(f"ИТОГО сумма: {total_sum:,.2f}".replace(",", " "))
    tr.bold = True
    tr.font.size = Pt(12)

    foot = doc.add_paragraph(
        f'Сформировано ИИ {dt.datetime.now().strftime("%d.%m.%Y %H:%M")} · Электронная Фармация')
    foot.runs[0].italic = True
    foot.runs[0].font.size = Pt(8)
    foot.runs[0].font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def build_xlsx(meta, cols, rows) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = "Отчёт"
    ws.append([meta["title"]])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f'Период: {meta["date_from"]:%d.%m.%Y} — {meta["date_to"]:%d.%m.%Y}'])
    ws.append([])
    ws.append(cols)
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append([(_num(v) if cols[i] in ("Сумма", "Цена", "Кол-во") else v) for i, v in enumerate(r)])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def is_marketing_prompt(prompt: str) -> bool:
    p = (prompt or "").lower()
    return any(k in p for k in MARKETING_KEYWORDS)


def marketing_text(prompt: str) -> str:
    """Рекламный текст про ЭльФиСА (не из данных). Best-effort через Ollama."""
    system = (
        "Ты — маркетолог платформы «ЭльФиСА» (Электронная Фармация): сервис заказа "
        "лекарств для аптек. Возможности: единый прайс от многих поставщиков в одном окне, "
        "быстрый поиск по названию, заказ в пару кликов, кросс-платформенный клиент "
        "(Windows/Mac/Linux) с авто-обновлением, ИИ-отчёты, выгрузка заказов поставщикам "
        "по их шаблонам (DBF/Excel/XML), справочники и регионы. "
        "Напиши бодрый, но правдоподобный рекламный текст (6-9 предложений) о том, почему "
        "ЭльФиСА — отличный выбор для аптеки. НЕ выдумывай конкретные цифры и статистику, "
        "не приписывай конкурентам ложных фактов. Пиши по-русски."
    )
    try:
        payload = {"model": OLLAMA_MODEL, "stream": False, "options": {"temperature": 0.8},
                   "messages": [{"role": "system", "content": system},
                                {"role": "user", "content": prompt}]}
        r = requests.post(f"{OLLAMA_URL}/api/chat", json=payload, timeout=HTTP_TIMEOUT)
        r.raise_for_status()
        return r.json().get("message", {}).get("content", "").strip()
    except Exception:
        return (
            "ЭльФиСА собирает прайсы многих поставщиков в одном окне — не нужно держать "
            "десяток программ и сайтов. Поиск по названию находит препарат за секунды, "
            "а заказ оформляется в пару кликов. Клиент работает на Windows, Mac и Linux и "
            "обновляется сам. ИИ-отчёты и удобная выгрузка заказов экономят время каждый день. "
            "Это простой и современный инструмент для ежедневной работы аптеки."
        )


def build_marketing_docx(prompt: str, text: str) -> bytes:
    doc = Document()
    h = doc.add_heading("Почему ЭльФиСА — отличный выбор", level=1)
    h.alignment = WD_ALIGN_PARAGRAPH.LEFT
    note = doc.add_paragraph(
        "Рекламный текст, сгенерирован ИИ по вашему запросу. Это не аналитика и не "
        "основано на данных ваших заказов.")
    note.runs[0].italic = True
    note.runs[0].font.color.rgb = RGBColor(0x88, 0x88, 0x88)
    for para in (text or "").split("\n"):
        para = para.strip()
        if para:
            doc.add_paragraph(para)
    foot = doc.add_paragraph(
        f'Сформировано ИИ {dt.datetime.now().strftime("%d.%m.%Y %H:%M")} · Электронная Фармация')
    foot.runs[0].italic = True
    foot.runs[0].font.size = Pt(8)
    foot.runs[0].font.color.rgb = RGBColor(0x88, 0x88, 0x88)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@app.get("/health")
def health():
    ok_ollama = False
    try:
        ok_ollama = requests.get(f"{OLLAMA_URL}/api/tags", timeout=5).ok
    except Exception:
        pass
    return {"status": "ok", "ollama": ok_ollama, "model": OLLAMA_MODEL, "api_base": API_BASE}


@app.post("/generate")
def generate(req: GenReq):
    if not req.token:
        raise HTTPException(status_code=400, detail="Нет токена")
    if not req.prompt or not req.prompt.strip():
        raise HTTPException(status_code=400, detail="Пустой запрос")

    # Маркетинговый/сравнительный запрос — не отчёт по данным, а рекламный текст.
    if is_marketing_prompt(req.prompt):
        text = marketing_text(req.prompt)
        data = build_marketing_docx(req.prompt, text)
        from urllib.parse import quote
        fname = "Почему ЭльФиСА.docx"
        headers = {
            "Content-Disposition": f"attachment; filename*=UTF-8''{quote(fname)}",
            "X-Report-Type": "marketing",
        }
        return StreamingResponse(
            io.BytesIO(data),
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers=headers)

    raw = extract_intent(req.prompt)
    meta = normalize_intent(raw, req.format)

    lines = fetch_lines(req.token, meta["date_from"], meta["date_to"])
    lines = apply_filters(lines, meta["supplier_filter"], meta["item_filter"])
    cols, rows = aggregate(lines, meta["report_type"])

    total_sum = sum(_num(l.get("sum")) for l in lines)
    period = f'{meta["date_from"]:%d.%m.%Y} — {meta["date_to"]:%d.%m.%Y}'
    summary = ai_summary(meta["title"], period, cols, rows, total_sum) if rows else ""

    safe = re.sub(r"[^\w\-. ]", "_", meta["title"])[:60].strip() or "report"
    if meta["format"] == "xlsx":
        data = build_xlsx(meta, cols, rows)
        media = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        fname = f"{safe}.xlsx"
    else:
        data = build_docx(meta, cols, rows, summary)
        media = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        fname = f"{safe}.docx"

    from urllib.parse import quote
    headers = {
        "Content-Disposition": f"attachment; filename*=UTF-8''{quote(fname)}",
        "X-Report-Type": meta["report_type"],
        "X-Report-Rows": str(len(rows)),
    }
    return StreamingResponse(io.BytesIO(data), media_type=media, headers=headers)


@app.post("/chat")
def chat(req: ChatReq):
    """Чат с ExestAI — общий ассистент кабинета (раздел «Чат»)."""
    if not req.message or not req.message.strip():
        raise HTTPException(status_code=400, detail="Пустое сообщение")
    messages = [{"role": "system", "content": EXEST_AI_SYSTEM}]
    for m in (req.history or [])[-10:]:
        role = m.get("role")
        content = (m.get("content") or "").strip()
        if role in ("user", "assistant") and content:
            messages.append({"role": role, "content": content})
    messages.append({"role": "user", "content": req.message.strip()})
    payload = {"model": OLLAMA_MODEL, "stream": False,
               "options": {"temperature": 0.4}, "messages": messages}
    try:
        r = requests.post(f"{OLLAMA_URL}/api/chat", json=payload, timeout=HTTP_TIMEOUT)
        r.raise_for_status()
        reply = r.json().get("message", {}).get("content", "").strip()
    except Exception as e:
        raise HTTPException(status_code=502, detail=f"ИИ недоступен: {e}")
    return JSONResponse({"reply": reply or "Извините, не смог сформировать ответ."})


@app.post("/map-template")
def map_template(req: MapReq):
    """
    Сопоставляет колонки шаблона накладной с полями системы.
    ИИ работает ОДИН РАЗ при загрузке шаблона; дальше выгрузка детерминированная.
    Вход: columns (колонки шаблона + примеры), fields (наши поля + описания).
    Выход: {"mappings": [{"column": <имя>, "field": <key|none>}]}.
    """
    field_lines = "\n".join(
        f'- {f.get("key")}: {f.get("description", "")}' for f in req.fields)
    col_lines = "\n".join(
        f'{i+1}. "{c.get("name")}" — примеры значений: {", ".join(str(s) for s in (c.get("samples") or [])[:8]) or "(пусто)"}'
        for i, c in enumerate(req.columns))

    system = (
        "Ты — эксперт по сопоставлению колонок накладной/заказа с полями системы. "
        "Для КАЖДОЙ колонки шаблона выбери ровно один field.key, который она означает, "
        "либо \"none\", если подходящего поля нет.\n"
        "Решай по ДВУМ признакам ВМЕСТЕ: (1) смысл заголовка колонки и (2) ЧТО РЕАЛЬНО лежит в примерах значений. "
        "Часто заголовок непонятный или на латинице — тогда именно по значениям определяй поле. "
        "Разбирай паттерны значений:\n"
        "- строка 8-4-4-4-12 hex (напр. 56cefd3c-c3fe-43ba-8c40-dcf62ccb0d20) — это GUID; по смыслу заголовка реши, какой именно id "
        "(order_id, order_item_id, price_list_id, price_list_item_id, goods_guid);\n"
        "- дата (2017-02-28 или 28.02.2017) → *_date; время (13:09:38) → order_time;\n"
        "- число с дробной частью (63.2000, 131.25, 2424.70) → цена (unit_price) или сумма (sum / order_total) по смыслу;\n"
        "- ровно 13 цифр (4603933014470) → штрихкод (barcode);\n"
        "- небольшие целые (1, 2, 3, 5) → количество (qty);\n"
        "- везде 0 или 0.0000 → это константа: qty_box / quantity? нет — qty_sheaf / vat_rate / alt_price по смыслу заголовка, иначе none;\n"
        "- строка вида Zkz-4546-18 → номер заказа (order_number);\n"
        "- название препарата/товара → item_name; название завода/фирмы → manufacturer; страна → country.\n"
        "Не назначай один и тот же field.key разным по значениям колонкам. "
        "Если сомневаешься и ни заголовок, ни значения не подходят — ставь \"none\".\n"
        "Верни СТРОГО JSON без пояснений: "
        "{\"mappings\": [{\"column\": \"<имя колонки>\", \"field\": \"<key или none>\"}]}."
    )
    user = f"ПОЛЯ СИСТЕМЫ:\n{field_lines}\n\nКОЛОНКИ ШАБЛОНА:\n{col_lines}"

    payload = {
        "model": OLLAMA_MODEL,
        "format": "json",
        "stream": False,
        "options": {"temperature": 0.1},
        "messages": [
            {"role": "system", "content": system},
            {"role": "user", "content": user},
        ],
    }
    r = requests.post(f"{OLLAMA_URL}/api/chat", json=payload, timeout=HTTP_TIMEOUT)
    r.raise_for_status()
    content = r.json().get("message", {}).get("content", "{}")
    try:
        return json.loads(content)
    except Exception:
        m = re.search(r"\{.*\}", content, re.S)
        return json.loads(m.group(0)) if m else {"mappings": []}
